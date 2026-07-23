from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
from bs4 import BeautifulSoup, Tag
from docx import Document as DocxDocument

from curx.domain.ingestion_schemas import SourceType

PAGE_MARKER_PATTERN = re.compile(r"<!--\s*curx-page:(\d+)\s*-->")
UNSAFE_BLOCK_PATTERN = re.compile(
    r"<(script|style|iframe|object|embed)\b[^>]*>.*?</\1\s*>",
    re.IGNORECASE | re.DOTALL,
)


class ContentLoadError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedDocument:
    markdown: str
    language: str
    metadata: dict[str, Any] = field(default_factory=dict)


def detect_source_type(filename: str, content_type: str | None) -> SourceType:
    suffix = Path(filename).suffix.lower()
    by_suffix = {
        ".md": SourceType.MARKDOWN,
        ".markdown": SourceType.MARKDOWN,
        ".html": SourceType.HTML,
        ".htm": SourceType.HTML,
        ".pdf": SourceType.PDF,
        ".docx": SourceType.DOCX,
    }
    if suffix in by_suffix:
        return by_suffix[suffix]

    normalized_type = (content_type or "").split(";", 1)[0].strip().lower()
    by_mime = {
        "text/markdown": SourceType.MARKDOWN,
        "text/x-markdown": SourceType.MARKDOWN,
        "text/html": SourceType.HTML,
        "application/pdf": SourceType.PDF,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
            SourceType.DOCX
        ),
    }
    try:
        return by_mime[normalized_type]
    except KeyError as exc:
        raise ContentLoadError(
            "unsupported_file_type",
            "Only Markdown, HTML, PDF and DOCX files are supported.",
        ) from exc


def normalize_markdown(value: str) -> str:
    sanitized = UNSAFE_BLOCK_PATTERN.sub("", value)
    sanitized = sanitized.replace("\x00", "")
    sanitized = re.sub(r"\r\n?", "\n", sanitized)
    sanitized = re.sub(r"[ \t]+\n", "\n", sanitized)
    sanitized = re.sub(r"\n{4,}", "\n\n\n", sanitized)
    normalized = sanitized.strip()
    if not normalized:
        raise ContentLoadError("empty_document", "The document contains no indexable text.")
    return normalized


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ContentLoadError(
        "invalid_text_encoding",
        "Text sources must use UTF-8 encoding.",
    )


def _language_for(text: str) -> str:
    letters = [character for character in text if character.isalpha()]
    if not letters:
        return "und"
    chinese = sum("\u4e00" <= character <= "\u9fff" for character in letters)
    return "zh" if chinese / len(letters) >= 0.2 else "en"


def _table_markdown(table: Tag) -> str:
    rows: list[list[str]] = []
    for row in table.find_all("tr"):
        cells = [
            re.sub(r"\s+", " ", cell.get_text(" ", strip=True)).replace("|", "\\|")
            for cell in row.find_all(["th", "td"])
        ]
        if cells:
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    body = rows[1:]
    lines = [
        f"| {' | '.join(header)} |",
        f"| {' | '.join(['---'] * width)} |",
    ]
    lines.extend(f"| {' | '.join(row)} |" for row in body)
    return "\n".join(lines)


def _html_to_markdown(content: bytes) -> ExtractedDocument:
    text = _decode_text(content)
    soup = BeautifulSoup(text, "html.parser")
    for unsafe in soup.find_all(["script", "style", "iframe", "object", "embed", "noscript"]):
        unsafe.decompose()

    blocks: list[str] = []
    root = soup.body or soup
    for element in root.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p", "ul", "ol", "table", "pre", "blockquote"],
        recursive=True,
    ):
        if element.find_parent(["p", "ul", "ol", "table", "pre", "blockquote"]):
            continue
        if element.name and element.name.startswith("h"):
            level = int(element.name[1])
            value = element.get_text(" ", strip=True)
            if value:
                blocks.append(f"{'#' * level} {value}")
        elif element.name == "p":
            value = element.get_text(" ", strip=True)
            if value:
                blocks.append(value)
        elif element.name in {"ul", "ol"}:
            lines = []
            for index, item in enumerate(element.find_all("li", recursive=False), start=1):
                marker = f"{index}." if element.name == "ol" else "-"
                lines.append(f"{marker} {item.get_text(' ', strip=True)}")
            if lines:
                blocks.append("\n".join(lines))
        elif element.name == "table":
            table = _table_markdown(element)
            if table:
                blocks.append(table)
        elif element.name == "pre":
            blocks.append(f"```\n{element.get_text().strip()}\n```")
        elif element.name == "blockquote":
            lines = element.get_text("\n", strip=True).splitlines()
            blocks.append("\n".join(f"> {line}" for line in lines if line.strip()))

    markdown = normalize_markdown("\n\n".join(blocks))
    html_title = (
        soup.title.string.strip()
        if soup.title and soup.title.string
        else None
    )
    return ExtractedDocument(
        markdown=markdown,
        language=_language_for(markdown),
        metadata={"html_title": html_title},
    )


def _pdf_to_markdown(content: bytes) -> ExtractedDocument:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ContentLoadError("damaged_pdf", "The PDF cannot be opened.") from exc

    try:
        pages: list[str] = []
        for page_index, page in enumerate(document, start=1):
            text = page.get_text("text", sort=True).strip()
            if not text:
                continue
            pages.append(
                f"<!-- curx-page:{page_index} -->\n\n"
                f"## 第 {page_index} 页\n\n{text}"
            )
        markdown = normalize_markdown("\n\n".join(pages))
        return ExtractedDocument(
            markdown=markdown,
            language=_language_for(markdown),
            metadata={"page_count": document.page_count},
        )
    finally:
        document.close()


def _docx_to_markdown(content: bytes) -> ExtractedDocument:
    try:
        document = DocxDocument(BytesIO(content))
    except Exception as exc:
        raise ContentLoadError("damaged_docx", "The DOCX file cannot be opened.") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        heading_match = re.match(r"Heading\s+([1-6])", style_name, re.IGNORECASE)
        if heading_match:
            blocks.append(f"{'#' * int(heading_match.group(1))} {value}")
        elif style_name.lower().startswith("list"):
            blocks.append(f"- {value}")
        else:
            blocks.append(value)

    for table in document.tables:
        rows = [
            [
                re.sub(r"\s+", " ", cell.text.strip()).replace("|", "\\|")
                for cell in row.cells
            ]
            for row in table.rows
        ]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        blocks.append(
            "\n".join(
                [
                    f"| {' | '.join(rows[0])} |",
                    f"| {' | '.join(['---'] * width)} |",
                    *(f"| {' | '.join(row)} |" for row in rows[1:]),
                ]
            )
        )

    markdown = normalize_markdown("\n\n".join(blocks))
    return ExtractedDocument(
        markdown=markdown,
        language=_language_for(markdown),
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )


def load_content(source_type: SourceType, content: bytes) -> ExtractedDocument:
    if source_type in {SourceType.MARKDOWN, SourceType.FAQ}:
        markdown = normalize_markdown(_decode_text(content))
        return ExtractedDocument(markdown=markdown, language=_language_for(markdown))
    if source_type is SourceType.HTML:
        return _html_to_markdown(content)
    if source_type is SourceType.PDF:
        return _pdf_to_markdown(content)
    if source_type is SourceType.DOCX:
        return _docx_to_markdown(content)
    raise ContentLoadError("unsupported_file_type", f"Unsupported source type: {source_type}.")
