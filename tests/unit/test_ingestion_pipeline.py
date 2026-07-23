from collections.abc import Iterator
from dataclasses import dataclass
from datetime import timedelta
from io import BytesIO
from uuid import UUID

import fitz
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from curx.core.config import Settings
from curx.core.database import Base, create_session_factory
from curx.domain.ingestion_schemas import IngestionStatus, SourceType
from curx.domain.models import (
    Chunk,
    ChunkEmbedding,
    Document,
    EmbeddingProfile,
    Source,
)
from curx.main import create_app
from curx.services.chunking import HierarchicalChunker
from curx.services.content_loaders import load_content
from curx.services.embeddings import LocalHashEmbedding
from curx.services.identity_service import bootstrap_admin, utc_now
from curx.services.job_queue import RecordingIngestionJobQueue
from curx.services.object_store import MemoryObjectStore
from curx.workers.ingestion import run_ingestion_job


@dataclass
class IngestionEnvironment:
    client: TestClient
    session_factory: sessionmaker[Session]
    settings: Settings
    object_store: MemoryObjectStore
    queue: RecordingIngestionJobQueue
    space_id: UUID

    def run(self, job_id: UUID) -> None:
        run_ingestion_job(
            job_id,
            session_factory=self.session_factory,
            object_store=self.object_store,
            embedding_model=LocalHashEmbedding(384, "hash-embedding-v1"),
            settings=self.settings,
        )


@pytest.fixture
def ingestion_environment() -> Iterator[IngestionEnvironment]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = create_session_factory(engine)
    settings = Settings(
        database_url="sqlite+pysqlite:///:memory:",
        session_cookie_secure=False,
        ingestion_inline=False,
        chunk_max_tokens=96,
        chunk_overlap_tokens=12,
        _env_file=None,
    )
    object_store = MemoryObjectStore()
    queue = RecordingIngestionJobQueue()
    with session_factory() as db:
        _, _, _, admin_key = bootstrap_admin(
            db,
            tenant_name="Curx Ingestion",
            tenant_slug="curx-ingestion",
            username="admin",
            display_name="Ingestion Admin",
            expires_in_days=30,
        )

    with TestClient(
        create_app(
            settings,
            session_factory,
            object_store=object_store,
            ingestion_job_queue=queue,
        )
    ) as client:
        bound = client.post(
            "/api/auth/bind",
            json={"username": "admin", "access_key": admin_key},
        )
        assert bound.status_code == 200
        workspace = client.get("/api/workspace").json()
        yield IngestionEnvironment(
            client=client,
            session_factory=session_factory,
            settings=settings,
            object_store=object_store,
            queue=queue,
            space_id=UUID(workspace["spaces"][0]["id"]),
        )
    engine.dispose()


def submit_markdown(environment: IngestionEnvironment, content: str) -> dict:
    response = environment.client.post(
        f"/api/knowledge-spaces/{environment.space_id}/sources/text",
        json={
            "title": "充电规范",
            "source_type": "markdown",
            "content": content,
            "tags": ["battery", "battery"],
        },
    )
    assert response.status_code == 202
    return response.json()


def test_markdown_submission_indexes_and_deduplicates(
    ingestion_environment: IngestionEnvironment,
) -> None:
    content = (
        "# 电池规范\n\n"
        "## 充电截止电压\n\n"
        "标准充电截止电压为 4.20V，允许误差为 ±0.05V。\n\n"
        "## 保护要求\n\n"
        "超过阈值时必须停止充电并记录故障。"
    )
    submitted = submit_markdown(ingestion_environment, content)
    job_id = UUID(submitted["job"]["id"])
    source_id = UUID(submitted["source"]["id"])
    assert submitted["job"]["status"] == "queued"
    assert submitted["source"]["tags"] == ["battery"]
    assert ingestion_environment.queue.job_ids == [job_id]

    ingestion_environment.run(job_id)

    job = ingestion_environment.client.get(f"/api/ingestion-jobs/{job_id}")
    assert job.status_code == 200
    assert job.json()["status"] == "indexed"
    assert [event["status"] for event in job.json()["events"]] == [
        "queued",
        "extracting",
        "normalizing",
        "chunking",
        "embedding",
        "indexed",
    ]
    detail = ingestion_environment.client.get(f"/api/sources/{source_id}")
    assert detail.status_code == 200
    assert detail.json()["documents"][0]["is_active"] is True
    assert detail.json()["documents"][0]["chunk_count"] >= 2

    duplicate = ingestion_environment.client.post(
        f"/api/knowledge-spaces/{ingestion_environment.space_id}/sources/text",
        json={
            "title": "同一份内容的新标题",
            "source_type": "markdown",
            "content": content,
        },
    )
    assert duplicate.status_code == 200
    assert duplicate.json()["deduplicated"] is True
    assert UUID(duplicate.json()["source"]["id"]) == source_id
    assert ingestion_environment.queue.job_ids == [job_id]

    with ingestion_environment.session_factory() as db:
        assert db.scalar(select(func.count(Source.id))) == 1
        assert db.scalar(select(func.count(Document.id))) == 1
        assert db.scalar(select(func.count(Chunk.id))) >= 2
        assert db.scalar(select(func.count(ChunkEmbedding.id))) == db.scalar(
            select(func.count(Chunk.id))
        )
        profile = db.scalar(select(EmbeddingProfile))
        assert profile is not None
        assert profile.provider == "local"
        assert profile.dimension == 384


def test_reindex_supersedes_old_version_without_duplicate_active_chunks(
    ingestion_environment: IngestionEnvironment,
) -> None:
    submitted = submit_markdown(
        ingestion_environment,
        "# 产品\n\n## 参数\n\n额定输入电压为 24V。",
    )
    first_job_id = UUID(submitted["job"]["id"])
    source_id = UUID(submitted["source"]["id"])
    ingestion_environment.run(first_job_id)

    reindex = ingestion_environment.client.post(f"/api/sources/{source_id}/reindex")
    assert reindex.status_code == 202
    second_job_id = UUID(reindex.json()["job"]["id"])
    ingestion_environment.run(second_job_id)

    with ingestion_environment.session_factory() as db:
        documents = list(
            db.scalars(
                select(Document)
                .where(Document.source_id == source_id)
                .order_by(Document.version_number)
            )
        )
        assert [document.status for document in documents] == ["superseded", "indexed"]
        assert [document.is_active for document in documents] == [False, True]
        active_chunks = db.scalar(
            select(func.count(Chunk.id)).where(
                Chunk.source_id == source_id,
                Chunk.is_active.is_(True),
            )
        )
        latest_chunks = db.scalar(
            select(func.count(Chunk.id)).where(Chunk.document_id == documents[1].id)
        )
        assert active_chunks == latest_chunks


def test_cancel_then_retry_runs_the_same_version(
    ingestion_environment: IngestionEnvironment,
) -> None:
    submitted = submit_markdown(
        ingestion_environment,
        "# 取消测试\n\n这份文档会先取消，再重试。",
    )
    job_id = UUID(submitted["job"]["id"])

    cancelled = ingestion_environment.client.post(
        f"/api/ingestion-jobs/{job_id}/cancel"
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"

    retried = ingestion_environment.client.post(f"/api/ingestion-jobs/{job_id}/retry")
    assert retried.status_code == 200
    assert retried.json()["status"] == "queued"
    ingestion_environment.run(job_id)
    completed = ingestion_environment.client.get(f"/api/ingestion-jobs/{job_id}").json()
    assert completed["status"] == "indexed"
    assert completed["attempt_count"] == 1


def test_unsupported_and_damaged_files_have_explicit_failures(
    ingestion_environment: IngestionEnvironment,
) -> None:
    unsupported = ingestion_environment.client.post(
        f"/api/knowledge-spaces/{ingestion_environment.space_id}/sources/upload",
        files={"file": ("malware.exe", b"MZ-not-really", "application/octet-stream")},
    )
    assert unsupported.status_code == 422
    assert unsupported.json()["detail"]["code"] == "unsupported_file_type"

    damaged = ingestion_environment.client.post(
        f"/api/knowledge-spaces/{ingestion_environment.space_id}/sources/upload",
        files={"file": ("broken.pdf", b"not a pdf", "application/pdf")},
    )
    assert damaged.status_code == 202
    job_id = UUID(damaged.json()["job"]["id"])
    ingestion_environment.run(job_id)
    failed = ingestion_environment.client.get(f"/api/ingestion-jobs/{job_id}").json()
    assert failed["status"] == "failed"
    assert failed["error_code"] == "damaged_pdf"


def test_member_can_list_but_cannot_upload_sources(
    ingestion_environment: IngestionEnvironment,
) -> None:
    created = ingestion_environment.client.post(
        "/api/admin/access-keys",
        json={
            "label": "Read only member",
            "granted_role": "member",
            "expires_at": (utc_now() + timedelta(days=7)).isoformat(),
            "knowledge_space_ids": [str(ingestion_environment.space_id)],
        },
    )
    assert created.status_code == 201

    with TestClient(ingestion_environment.client.app) as member_client:
        bound = member_client.post(
            "/api/auth/bind",
            json={
                "username": "reader",
                "display_name": "Reader",
                "access_key": created.json()["access_key"],
            },
        )
        assert bound.status_code == 200
        listed = member_client.get(
            f"/api/knowledge-spaces/{ingestion_environment.space_id}/sources"
        )
        assert listed.status_code == 200
        denied = member_client.post(
            f"/api/knowledge-spaces/{ingestion_environment.space_id}/sources/text",
            json={
                "title": "Forbidden",
                "source_type": "markdown",
                "content": "# Secret",
            },
        )
        assert denied.status_code == 403
        assert denied.json()["detail"]["code"] == "knowledge_write_required"


def test_html_pdf_docx_and_hierarchical_chunk_locations() -> None:
    html = load_content(
        source_type=SourceType.HTML,
        content=(
            b"<html><head><title>Guide</title><script>bad()</script></head>"
            b"<body><h1>Setup</h1><p>Use 24V.</p>"
            b"<table><tr><th>Mode</th><th>Value</th></tr>"
            b"<tr><td>Safe</td><td>24V</td></tr></table></body></html>"
        ),
    )
    assert "# Setup" in html.markdown
    assert "| Mode | Value |" in html.markdown
    assert "bad()" not in html.markdown

    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "Battery cutoff voltage is 4.20V.")
    pdf_bytes = pdf_document.tobytes()
    pdf_document.close()
    pdf = load_content(source_type=SourceType.PDF, content=pdf_bytes)
    assert "<!-- curx-page:1 -->" in pdf.markdown

    docx_document = DocxDocument()
    docx_document.add_heading("Safety", level=1)
    docx_document.add_paragraph("Disconnect power before maintenance.")
    docx_buffer = BytesIO()
    docx_document.save(docx_buffer)
    docx = load_content(source_type=SourceType.DOCX, content=docx_buffer.getvalue())
    assert "# Safety" in docx.markdown

    chunks = HierarchicalChunker(max_tokens=96, overlap_tokens=12).chunk(pdf.markdown)
    assert chunks[0].page_number == 1
    assert "第 1 页" in chunks[0].location


def test_job_state_enum_remains_closed() -> None:
    assert {status.value for status in IngestionStatus} == {
        "queued",
        "extracting",
        "normalizing",
        "chunking",
        "embedding",
        "indexed",
        "failed",
        "cancelled",
        "superseded",
    }
